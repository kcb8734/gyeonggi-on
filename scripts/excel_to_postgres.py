#!/usr/bin/env python3
"""엑셀(.xlsx) 데이터를 PostgreSQL에 적재합니다.

경기온(Gyeonggi-On) 마스터 시트(municipalities / festivals / merchants /
discount_promotions / coupons)는 한글·영문 헤더와 외래키 조회를 지원합니다.
그 외 시트는 테이블 컬럼명과 헤더가 일치하면 그대로 INSERT/UPSERT 합니다.

사용 예:

  pip install -r scripts/requirements-excel.txt

  # 템플릿 엑셀 생성
  python scripts/excel_to_postgres.py --create-template ./import_template.xlsx

  # 미리보기 (DB에 커밋하지 않음)
  python scripts/excel_to_postgres.py --file ./import_template.xlsx --dry-run

  # 전체 시트 적재 (backend/.env 의 DATABASE_URL 사용)
  python scripts/excel_to_postgres.py --file ./data.xlsx

  # 특정 시트만
  python scripts/excel_to_postgres.py --file ./data.xlsx --sheet 가맹점 --table merchants
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import urlparse, urlunparse

try:
    from dotenv import load_dotenv
except ImportError:  # pragma: no cover
    load_dotenv = None

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent
BACKEND_DIR = REPO_ROOT / "backend"

TRUTHY = {"1", "true", "t", "y", "yes", "o", "예", "맞음", "on", "활성"}
FALSY = {"0", "false", "f", "n", "no", "x", "아니오", "아님", "off", "비활성"}

SKIP_SHEETS = {
    "총괄": "요약 시트",
    "응답보기": "문항/코드 시트",
    "안내": "안내 시트",
    "표지": "표지 시트",
    "코드북": "문항/코드 시트",
}

SHEET_ALIASES = {
    "municipalities": "municipalities",
    "지자체": "municipalities",
    "시군구": "municipalities",
    "festivals": "festivals",
    "축제": "festivals",
    "축제정보": "festivals",
    "조사표": "festivals",
    "개최계획": "festivals",
    "개최현황": "festivals",
    "지역축제": "festivals",
    "축제현황": "festivals",
    "문화축제": "festivals",
    "merchants": "merchants",
    "가맹점": "merchants",
    "점포": "merchants",
    "소상공인": "merchants",
    "discount_promotions": "discount_promotions",
    "promotions": "discount_promotions",
    "프로모션": "discount_promotions",
    "할인": "discount_promotions",
    "coupons": "coupons",
    "쿠폰": "coupons",
}

LOAD_ORDER = [
    "municipalities",
    "festivals",
    "merchants",
    "discount_promotions",
    "coupons",
]


def canon(value: Any) -> str:
    text = "" if value is None else str(value).strip().lower()
    return re.sub(r"[\s_\-()/]+", "", text)


def is_skipped_sheet(sheet_name: str) -> bool:
    key = canon(sheet_name)
    return any(canon(alias) == key or key.startswith(canon(alias)) for alias in SKIP_SHEETS)


def normalize_header(value: Any) -> str:
    text = "" if value is None else str(value).strip()
    text = re.sub(r"^\s*(?:문항|질문|항목)?\s*q?\s*\d+\s*[\.\)\:\-]", "", text, flags=re.I)
    text = re.sub(r"[\*＊]+", "", text)
    text = re.sub(r"\([^)]*\)", "", text)
    return re.sub(r"\s+", " ", text).strip()


def is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    try:
        import math

        if isinstance(value, float) and math.isnan(value):
            return True
    except Exception:
        pass
    return False


def to_bool(value: Any, default: bool | None = None) -> bool | None:
    if is_blank(value):
        return default
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return bool(value)
    key = canon(value)
    if key in TRUTHY:
        return True
    if key in FALSY:
        return False
    raise ValueError(f"불리언으로 변환할 수 없습니다: {value!r}")


def to_number(value: Any) -> Decimal | None:
    if is_blank(value):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        raise ValueError("불리언은 숫자로 쓰지 않습니다")
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip().replace(",", "")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"숫자로 변환할 수 없습니다: {value!r}") from exc


def to_int(value: Any) -> int | None:
    number = to_number(value)
    if number is None:
        return None
    return int(number)


def to_text(value: Any) -> str | None:
    if is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def to_date(value: Any, year_hint: int | None = None) -> date | None:
    if is_blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and 59 < float(value) < 80000:
        return date(1899, 12, 30) + timedelta(days=int(value))
    text = str(value).strip()
    full = re.search(r"(20\d{2})\s*[.\-/년]?\s*(\d{1,2})\s*[.\-/월]?\s*(\d{1,2})", text)
    if full:
        return date(int(full.group(1)), int(full.group(2)), int(full.group(3)))
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(text[:10] if fmt != "%Y%m%d" else text[:8], fmt).date()
        except ValueError:
            continue
    short = re.match(r"^(\d{1,2})[.\-/](\d{1,2})$", text)
    if short:
        year = year_hint or date.today().year
        return date(year, int(short.group(1)), int(short.group(2)))
    raise ValueError(f"날짜로 변환할 수 없습니다: {value!r}")


def parse_period(value: Any, year_hint: int | None = None) -> tuple[date | None, date | None]:
    if is_blank(value):
        return (None, None)
    if isinstance(value, (date, datetime)):
        parsed = to_date(value)
        return (parsed, parsed)
    text = str(value).strip()
    full = list(re.finditer(r"(20\d{2})\s*[.\-/년]\s*(\d{1,2})\s*[.\-/월]\s*(\d{1,2})", text))
    if full:
        start = date(int(full[0].group(1)), int(full[0].group(2)), int(full[0].group(3)))
        end = start
        if len(full) > 1:
            end = date(int(full[1].group(1)), int(full[1].group(2)), int(full[1].group(3)))
        return (start, end)
    shorts = list(re.finditer(r"(\d{1,2})\s*[.\-/월]\s*(\d{1,2})\s*일?", text))
    if shorts:
        year = year_hint or date.today().year
        start = date(year, int(shorts[0].group(1)), int(shorts[0].group(2)))
        end = start
        if len(shorts) > 1:
            end = date(year, int(shorts[1].group(1)), int(shorts[1].group(2)))
        return (start, end)
    try:
        parsed = to_date(value, year_hint)
        return (parsed, parsed)
    except ValueError:
        return (None, None)


def to_datetime(value: Any) -> datetime | None:
    if is_blank(value):
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    if isinstance(value, date) and not isinstance(value, datetime):
        return datetime.combine(value, time.min, tzinfo=timezone.utc)
    text = str(value).strip().replace("T", " ")
    for fmt in (
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d",
    ):
        try:
            parsed = datetime.strptime(text[:19], fmt)
            return parsed.replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    raise ValueError(f"일시로 변환할 수 없습니다: {value!r}")


def to_uuid(value: Any) -> str | None:
    text = to_text(value)
    if text is None:
        return None
    return str(uuid.UUID(text))


CONVERTERS: dict[str, Callable[[Any], Any]] = {
    "text": to_text,
    "bool": to_bool,
    "int": to_int,
    "number": to_number,
    "date": to_date,
    "datetime": to_datetime,
    "uuid": to_uuid,
}


@dataclass
class Column:
    name: str
    kind: str = "text"
    aliases: tuple[str, ...] = ()
    required: bool = False
    skip_insert: bool = False

    def keys(self) -> set[str]:
        return {canon(self.name), *(canon(alias) for alias in self.aliases)}


@dataclass
class TableProfile:
    table: str
    columns: list[Column]
    conflict: tuple[str, ...] = ()
    skip_update: tuple[str, ...] = ()
    defaults: dict[str, Any] = field(default_factory=dict)

    def by_alias(self) -> dict[str, Column]:
        mapping: dict[str, Column] = {}
        for column in self.columns:
            for key in column.keys():
                mapping[key] = column
        return mapping

    def insertable(self) -> list[Column]:
        return [column for column in self.columns if not column.skip_insert]


PROFILES: dict[str, TableProfile] = {
    "municipalities": TableProfile(
        table="municipalities",
        conflict=("region_code",),
        columns=[
            Column("id", "uuid", ("ID",)),
            Column("name", "text", ("지자체명", "시군구", "이름"), required=True),
            Column("region_code", "text", ("지역코드", "코드"), required=True),
            Column("metro_region", "text", ("권역", "광역")),
            Column("budget_balance", "number", ("예산잔액", "예산")),
            Column("initial_budget", "number", ("초기예산",)),
            Column("contact_email", "text", ("담당자이메일", "연락처이메일")),
            Column("settlement_email", "text", ("정산이메일",)),
            Column("mayor_name", "text", ("단체장", "시장", "군수")),
            Column("department", "text", ("담당부서", "부서")),
        ],
        defaults={"metro_region": "GYEONGGI", "budget_balance": Decimal("0")},
    ),
    "festivals": TableProfile(
        table="festivals",
        conflict=("tour_content_id",),
        skip_update=("id",),
        columns=[
            Column("id", "uuid", ("ID",)),
            Column("municipality_id", "uuid", ("지자체ID",), skip_insert=True),
            Column("title", "text", ("축제명", "행사명", "축제이름", "제목", "이름"), required=True),
            Column("description", "text", ("설명", "소개", "축제내용", "주요내용", "행사내용")),
            Column("start_date", "date", ("시작일", "시작일자", "개최시작일", "축제시작일자", "축제시작일", "시작", "개최기간시작"), required=True),
            Column("end_date", "date", ("종료일", "종료일자", "개최종료일", "축제종료일자", "축제종료일", "종료", "개최기간종료"), required=True),
            Column("location_name", "text", ("장소", "주소", "개최장소", "행사장소", "소재지도로명주소")),
            Column("latitude", "number", ("위도", "lat")),
            Column("longitude", "number", ("경도", "lng", "lon")),
            Column("category", "text", ("카테고리", "분류")),
            Column("image_url", "text", ("이미지", "이미지URL")),
            Column("is_trending", "bool", ("인기", "트렌딩")),
            Column("tour_content_id", "text", ("관광공사ID", "contentid", "content_id")),
            Column("tel", "text", ("전화번호", "연락처")),
            Column("source", "text", ("출처",)),
        ],
        defaults={"category": "문화/예술", "source": "excel", "is_trending": False},
    ),
    "merchants": TableProfile(
        table="merchants",
        conflict=("business_number",),
        skip_update=("id", "owner_user_id"),
        columns=[
            Column("id", "uuid", ("ID",)),
            Column("owner_user_id", "uuid", ("점주ID", "소유자ID")),
            Column("municipality_id", "uuid", ("지자체ID",), skip_insert=True),
            Column("business_name", "text", ("상호명", "상호", "점포명"), required=True),
            Column("business_number", "text", ("사업자등록번호", "사업자번호"), required=True),
            Column("category", "text", ("업종", "카테고리"), required=True),
            Column("address", "text", ("주소",), required=True),
            Column("latitude", "number", ("위도", "lat")),
            Column("longitude", "number", ("경도", "lng", "lon")),
            Column("bank_name", "text", ("은행", "은행명")),
            Column("bank_account_number", "text", ("계좌번호", "정산계좌")),
            Column("is_verified", "bool", ("인증여부", "인증")),
        ],
        defaults={"is_verified": False},
    ),
    "discount_promotions": TableProfile(
        table="discount_promotions",
        conflict=("id",),
        columns=[
            Column("id", "uuid", ("ID",)),
            Column("merchant_id", "uuid", ("가맹점ID",), skip_insert=True),
            Column("festival_id", "uuid", ("축제ID",), skip_insert=True),
            Column("title", "text", ("프로모션명", "제목"), required=True),
            Column("merchant_discount_rate", "number", ("점주할인율", "자체할인율"), required=True),
            Column("gov_matching_rate", "number", ("지자체매칭율", "매칭율")),
            Column("max_discount_amount", "number", ("최대할인액", "할인한도")),
            Column("total_quantity", "int", ("총수량", "발급수량"), required=True),
            Column("remaining_quantity", "int", ("잔여수량",)),
            Column("start_time", "datetime", ("시작일시", "시작"), required=True),
            Column("end_time", "datetime", ("종료일시", "종료"), required=True),
            Column("status", "text", ("상태",)),
            Column("funding_type", "text", ("재원유형",)),
            Column("matching_status", "text", ("매칭상태",)),
            Column("coupon_type", "text", ("쿠폰유형",)),
        ],
        defaults={
            "gov_matching_rate": Decimal("0"),
            "status": "ACTIVE",
            "funding_type": "MATCHED",
            "matching_status": "NONE",
            "coupon_type": "OFFICIAL",
        },
    ),
    "coupons": TableProfile(
        table="coupons",
        conflict=("code",),
        columns=[
            Column("id", "uuid", ("ID",)),
            Column("code", "text", ("쿠폰코드", "코드"), required=True),
            Column("title", "text", ("쿠폰명", "제목"), required=True),
            Column("discount_amount", "int", ("할인금액", "할인액"), required=True),
            Column("municipality_id", "uuid", ("지자체ID",), skip_insert=True),
            Column("merchant_id", "uuid", ("가맹점ID",), skip_insert=True),
            Column("is_used", "bool", ("사용여부",)),
            Column("used_at", "datetime", ("사용일시",)),
            Column("expires_at", "datetime", ("만료일시", "만료일"), required=True),
            Column("coupon_type", "text", ("쿠폰유형",)),
        ],
        defaults={"is_used": False, "coupon_type": "OFFICIAL"},
    ),
}


class LoadError(RuntimeError):
    pass


def load_env_files() -> list[str]:
    loaded: list[str] = []
    if load_dotenv is None:
        return loaded
    for path in (
        BACKEND_DIR / ".env",
        REPO_ROOT / ".env",
        Path.cwd() / ".env",
    ):
        if path.is_file():
            load_dotenv(path, override=False)
            loaded.append(str(path))
    return loaded


def database_url() -> str:
    return (os.environ.get("DATABASE_URL") or os.environ.get("POSTGRES_URL") or "").strip()


def database_host(url: str) -> str:
    if not url:
        return ""
    try:
        return urlparse(url).hostname or ""
    except Exception:
        matched = re.search(r"@([^/?]+)", url)
        return (matched.group(1).split(":")[0] if matched else "")


def should_use_ssl(url: str) -> bool:
    if not url or re.search(r"sslmode=disable", url, re.I):
        return False
    host = database_host(url)
    if host in {"localhost", "127.0.0.1", "::1"}:
        return False
    return "neon.tech" in host or "neon.cloud" in host or bool(re.search(r"sslmode=require", url, re.I)) or bool(host)


def connection_dsn(url: str) -> str:
    if not should_use_ssl(url):
        return url
    try:
        parsed = urlparse(url)
        query = "&".join(
            part for part in parsed.query.split("&") if part and not part.lower().startswith("sslmode=")
        )
        return urlunparse(parsed._replace(query=query))
    except Exception:
        return re.sub(r"([?&])sslmode=[^&]+", r"\1", url).rstrip("?&")


def resolve_table_name(sheet_name: str, forced: str | None = None) -> str:
    if forced:
        return forced
    if is_skipped_sheet(sheet_name):
        return ""
    key = canon(sheet_name)
    if key in SHEET_ALIASES:
        return SHEET_ALIASES[key]
    return sheet_name.strip()


def pick(row: dict[str, Any], *names: str) -> Any:
    by_key: dict[str, Any] = {}
    for key, value in row.items():
        by_key[canon(key)] = value
        normalized = normalize_header(key)
        if normalized:
            by_key[canon(normalized)] = value
    for name in names:
        key = canon(name)
        if key in by_key and not is_blank(by_key[key]):
            return by_key[key]
    return None


def apply_profile(row: dict[str, Any], profile: TableProfile) -> dict[str, Any]:
    source = dict(row)
    if profile.table == "festivals":
        start, end = parse_period(pick(source, "개최기간", "행사기간", "축제기간", "기간"))
        if start and is_blank(pick(source, "시작일", "축제시작일자", "축제시작일", "개최시작일")):
            source["시작일"] = start
        if end and is_blank(pick(source, "종료일", "축제종료일자", "축제종료일", "개최종료일")):
            source["종료일"] = end
        city = pick(source, "시군구명", "시군구", "시군", "기초단체", "지자체명")
        if city and is_blank(source.get("시군구")):
            source["시군구"] = city
    out: dict[str, Any] = {}
    raw_by_canon: dict[str, Any] = {}
    for key, value in source.items():
        raw_by_canon[canon(key)] = value
        normalized = normalize_header(key)
        if normalized:
            raw_by_canon[canon(normalized)] = value
    for column in profile.columns:
        raw = None
        for key in column.keys():
            if key in raw_by_canon:
                raw = raw_by_canon[key]
                break
        if is_blank(raw):
            if column.required and column.name in profile.defaults:
                out[column.name] = profile.defaults[column.name]
            continue
        converter = CONVERTERS[column.kind]
        out[column.name] = converter(raw)
    if profile.table == "merchants" and not out.get("owner_user_id"):
        out["owner_user_id"] = str(uuid.uuid4())
    if profile.table == "discount_promotions" and out.get("remaining_quantity") is None and out.get("total_quantity") is not None:
        out["remaining_quantity"] = out["total_quantity"]
    if profile.table == "festivals" and not out.get("end_date") and out.get("start_date"):
        out["end_date"] = out["start_date"]
    missing = [
        column.name
        for column in profile.columns
        if column.required and is_blank(out.get(column.name))
    ]
    if missing:
        raise LoadError(f"{profile.table} 필수 값이 없습니다: {', '.join(missing)}")
    return out


def read_excel_sheets(path: Path) -> list[tuple[str, list[str], list[dict[str, Any]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise LoadError("openpyxl 이 필요합니다. pip install -r scripts/requirements-excel.txt") from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    sheets: list[tuple[str, list[str], list[dict[str, Any]]]] = []
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            try:
                header_row = next(rows)
            except StopIteration:
                continue
            if not header_row or all(is_blank(cell) for cell in header_row):
                continue
            headers = [str(cell).strip() if not is_blank(cell) else "" for cell in header_row]
            records: list[dict[str, Any]] = []
            for values in rows:
                if not values or all(is_blank(cell) for cell in values):
                    continue
                record = {
                    headers[index]: values[index]
                    for index in range(min(len(headers), len(values)))
                    if headers[index]
                }
                records.append(record)
            sheets.append((worksheet.title, [h for h in headers if h], records))
    finally:
        workbook.close()
    return sheets


def create_template(path: Path) -> Path:
    from openpyxl import Workbook

    workbook = Workbook()
    sheets = {
        "지자체": [
            ["지자체명", "지역코드", "권역", "예산잔액", "담당자이메일"],
            ["수원시", "GG_SUWON", "GYEONGGI", 50000000, "suwon@example.go.kr"],
            ["용인시", "GG_YONGIN", "GYEONGGI", 30000000, "yongin@example.go.kr"],
        ],
        "축제": [
            ["축제명", "지자체명", "지역코드", "시작일", "종료일", "장소", "위도", "경도", "카테고리", "관광공사ID"],
            ["수원화성문화제", "수원시", "GG_SUWON", "2026-09-01", "2026-09-30", "수원화성 행궁광장", 37.2870, 127.0130, "문화/예술", "EXCEL-FEST-001"],
        ],
        "가맹점": [
            ["상호명", "사업자등록번호", "업종", "주소", "지자체명", "지역코드", "위도", "경도", "인증여부"],
            ["화성행궁 한정식", "123-45-00001", "음식점", "경기도 수원시 팔달구 정조로 825", "수원시", "GG_SUWON", 37.2865, 127.0135, "예"],
        ],
        "프로모션": [
            [
                "프로모션명",
                "사업자등록번호",
                "축제명",
                "점주할인율",
                "지자체매칭율",
                "총수량",
                "시작일시",
                "종료일시",
            ],
            [
                "화성문화제 제휴 10% 할인",
                "123-45-00001",
                "수원화성문화제",
                5,
                5,
                100,
                "2026-09-01 00:00:00",
                "2026-09-30 23:59:59",
            ],
        ],
    }
    first = True
    for title, rows in sheets.items():
        worksheet = workbook.active if first else workbook.create_sheet(title)
        if first:
            worksheet.title = title
            first = False
        for row in rows:
            worksheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


class PostgresLoader:
    def __init__(self, dsn: str, dry_run: bool = False, create_missing: bool = True):
        try:
            import psycopg2
            from psycopg2.extras import RealDictCursor
        except ImportError as exc:  # pragma: no cover
            raise LoadError("psycopg2 가 필요합니다. pip install -r scripts/requirements-excel.txt") from exc

        self.psycopg2 = psycopg2
        self.dry_run = dry_run
        self.create_missing = create_missing
        self.conn = psycopg2.connect(
            connection_dsn(dsn),
            sslmode="require" if should_use_ssl(dsn) else "prefer",
            cursor_factory=RealDictCursor,
        )
        self.conn.autocommit = False
        self._muni_cache: dict[str, str] = {}
        self._merchant_cache: dict[str, str] = {}
        self._festival_cache: dict[str, str] = {}

    def close(self) -> None:
        self.conn.close()

    def commit_or_rollback(self) -> None:
        if self.dry_run:
            self.conn.rollback()
        else:
            self.conn.commit()

    def rollback(self) -> None:
        self.conn.rollback()

    def table_columns(self, table: str) -> list[str]:
        with self.conn.cursor() as cur:
            cur.execute(
                """
                SELECT column_name
                FROM information_schema.columns
                WHERE table_schema = 'public' AND table_name = %s
                ORDER BY ordinal_position
                """,
                (table,),
            )
            return [row["column_name"] for row in cur.fetchall()]

    def _one_id(self, sql: str, params: tuple[Any, ...]) -> str | None:
        with self.conn.cursor() as cur:
            cur.execute(sql, params)
            row = cur.fetchone()
            return str(row["id"]) if row and row.get("id") else None

    def resolve_municipality_id(self, row: dict[str, Any], create: bool = False) -> str | None:
        if row.get("municipality_id"):
            return str(row["municipality_id"])
        name = to_text(pick(row, "municipality", "municipality_name", "지자체", "지자체명", "시군구", "name"))
        region = to_text(pick(row, "region_code", "지역코드"))
        if not name and not region:
            return None
        cache_key = canon(region or name)
        if cache_key in self._muni_cache:
            return self._muni_cache[cache_key]
        found = self._one_id(
            """
            SELECT id FROM municipalities
            WHERE (%s IS NOT NULL AND name = %s)
               OR (%s IS NOT NULL AND region_code = %s)
            LIMIT 1
            """,
            (name, name, region, region),
        )
        if found:
            self._muni_cache[cache_key] = found
            if name:
                self._muni_cache[canon(name)] = found
            if region:
                self._muni_cache[canon(region)] = found
            return found
        if not create:
            return None
        if not name:
            raise LoadError("지자체를 새로 만들려면 지자체명이 필요합니다")
        region = region or ("GG_" + re.sub(r"\s+", "", name))
        with self.conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO municipalities (name, region_code, budget_balance, metro_region)
                VALUES (%s, %s, 0, COALESCE(%s, 'GYEONGGI'))
                ON CONFLICT (region_code) DO UPDATE SET name = EXCLUDED.name
                RETURNING id
                """,
                (name, region, to_text(pick(row, "metro_region", "권역"))),
            )
            found = str(cur.fetchone()["id"])
        self._muni_cache[canon(name)] = found
        self._muni_cache[canon(region)] = found
        return found

    def resolve_merchant_id(self, row: dict[str, Any]) -> str | None:
        if row.get("merchant_id"):
            return str(row["merchant_id"])
        number = to_text(pick(row, "business_number", "사업자등록번호", "사업자번호"))
        name = to_text(pick(row, "merchant", "business_name", "상호명", "상호"))
        if number:
            key = canon(number)
            if key in self._merchant_cache:
                return self._merchant_cache[key]
            found = self._one_id("SELECT id FROM merchants WHERE business_number = %s LIMIT 1", (number,))
            if found:
                self._merchant_cache[key] = found
            return found
        if name:
            key = canon(name)
            if key in self._merchant_cache:
                return self._merchant_cache[key]
            found = self._one_id("SELECT id FROM merchants WHERE business_name = %s LIMIT 1", (name,))
            if found:
                self._merchant_cache[key] = found
            return found
        return None

    def resolve_festival_id(self, row: dict[str, Any]) -> str | None:
        if row.get("festival_id"):
            return str(row["festival_id"])
        content_id = to_text(pick(row, "tour_content_id", "관광공사ID", "contentid"))
        title = to_text(pick(row, "festival", "festival_title", "축제명", "축제"))
        if content_id:
            found = self._one_id("SELECT id FROM festivals WHERE tour_content_id = %s LIMIT 1", (content_id,))
            if found:
                return found
        if title:
            key = canon(title)
            if key in self._festival_cache:
                return self._festival_cache[key]
            found = self._one_id(
                "SELECT id FROM festivals WHERE title = %s ORDER BY created_at DESC LIMIT 1",
                (title,),
            )
            if found:
                self._festival_cache[key] = found
            return found
        return None

    def enrich(self, table: str, source: dict[str, Any], mapped: dict[str, Any]) -> dict[str, Any]:
        merged = dict(source)
        merged.update(mapped)
        if table in {"festivals", "merchants", "coupons"}:
            municipality_id = self.resolve_municipality_id(
                merged, create=self.create_missing and table in {"festivals", "merchants"}
            )
            if municipality_id:
                mapped["municipality_id"] = municipality_id
            elif table in {"festivals", "merchants"}:
                raise LoadError("지자체(지자체명 또는 지역코드)를 찾을 수 없습니다")
        if table in {"discount_promotions", "coupons"}:
            merchant_id = self.resolve_merchant_id(merged)
            if merchant_id:
                mapped["merchant_id"] = merchant_id
            elif table == "discount_promotions":
                raise LoadError("가맹점(사업자등록번호 또는 상호명)을 찾을 수 없습니다")
        if table == "discount_promotions":
            festival_id = self.resolve_festival_id(merged)
            if festival_id:
                mapped["festival_id"] = festival_id
        return mapped

    def upsert(self, table: str, row: dict[str, Any], conflict: tuple[str, ...], skip_update: tuple[str, ...] = ()) -> str:
        usable = {key: value for key, value in row.items() if value is not None}
        if not usable:
            raise LoadError("적재할 컬럼이 없습니다")
        cols = list(usable.keys())
        sql = build_upsert_sql(table, cols, conflict, skip_update)
        with self.conn.cursor() as cur:
            cur.execute(sql, [usable[col] for col in cols])
            fetched = cur.fetchone()
            if fetched and fetched.get("id"):
                return str(fetched["id"])
        return ""


def build_upsert_sql(
    table: str,
    cols: list[str],
    conflict: tuple[str, ...] | Iterable[str] = (),
    skip_update: tuple[str, ...] | Iterable[str] = (),
) -> str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", table):
        raise LoadError(f"허용되지 않는 테이블 이름입니다: {table}")
    for col in cols:
        if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", col):
            raise LoadError(f"허용되지 않는 컬럼 이름입니다: {col}")
    placeholders = ", ".join(["%s"] * len(cols))
    col_sql = ", ".join(cols)
    conflict_cols = [col for col in conflict if col in cols]
    sql = f"INSERT INTO {table} ({col_sql}) VALUES ({placeholders})"
    if conflict_cols:
        skip = set(skip_update)
        update_cols = [col for col in cols if col not in conflict_cols and col not in skip]
        if update_cols:
            assignments = ", ".join(f"{col} = EXCLUDED.{col}" for col in update_cols)
            sql += f" ON CONFLICT ({', '.join(conflict_cols)}) DO UPDATE SET {assignments}"
        else:
            sql += f" ON CONFLICT ({', '.join(conflict_cols)}) DO NOTHING"
    sql += " RETURNING id"
    return sql


def sheet_sort_key(sheet_name: str) -> tuple[int, str]:
    table = resolve_table_name(sheet_name)
    try:
        return (LOAD_ORDER.index(table), sheet_name)
    except ValueError:
        return (len(LOAD_ORDER), sheet_name)


def load_file(
    path: Path,
    *,
    sheet: str | None = None,
    table: str | None = None,
    dry_run: bool = False,
    create_missing: bool = True,
    dsn: str | None = None,
) -> dict[str, Any]:
    if not path.is_file():
        raise LoadError(f"엑셀 파일이 없습니다: {path}")
    url = (database_url() if dsn is None else dsn).strip()
    if not url:
        raise LoadError("DATABASE_URL 또는 POSTGRES_URL 이 필요합니다")

    sheets = read_excel_sheets(path)
    if sheet:
        sheets = [item for item in sheets if item[0] == sheet or canon(item[0]) == canon(sheet)]
        if not sheets:
            raise LoadError(f"시트를 찾을 수 없습니다: {sheet}")
    sheets.sort(key=lambda item: sheet_sort_key(item[0]))

    loader = PostgresLoader(url, dry_run=dry_run, create_missing=create_missing)
    summary: dict[str, Any] = {"file": str(path), "dry_run": dry_run, "sheets": []}
    try:
        for sheet_name, _headers, records in sheets:
            if is_skipped_sheet(sheet_name) and not table:
                summary["sheets"].append({
                    "sheet": sheet_name,
                    "table": None,
                    "inserted": 0,
                    "skipped": True,
                    "errors": [],
                })
                continue
            table_name = resolve_table_name(sheet_name, table)
            profile = PROFILES.get(table_name)
            if profile is None:
                db_cols = set(loader.table_columns(table_name))
                if not db_cols:
                    raise LoadError(f"알 수 없는 테이블입니다: {table_name}")
            result = {"sheet": sheet_name, "table": table_name, "inserted": 0, "errors": []}
            for index, raw in enumerate(records, start=2):
                try:
                    mapped = apply_profile(raw, profile) if profile else {
                        key: value for key, value in raw.items() if not is_blank(value)
                    }
                    mapped = loader.enrich(table_name, raw, mapped)
                    if profile:
                        mapped = {
                            key: value
                            for key, value in mapped.items()
                            if key in {col.name for col in profile.insertable()} or key.endswith("_id")
                        }
                        if table_name == "festivals" and not mapped.get("tour_content_id"):
                            conflict = ("id",) if mapped.get("id") else ()
                        elif table_name == "discount_promotions" and not mapped.get("id"):
                            conflict = ()
                        else:
                            conflict = profile.conflict
                        skip_update = profile.skip_update
                    else:
                        conflict = ()
                        skip_update = ()
                    loader.upsert(table_name, mapped, conflict, skip_update)
                    result["inserted"] += 1
                except Exception as exc:
                    result["errors"].append({"row": index, "error": str(exc)})
                    loader.rollback()
                    summary["sheets"].append(result)
                    raise LoadError(f"{sheet_name} {index}행 실패: {exc}") from exc
            summary["sheets"].append(result)
        loader.commit_or_rollback()
    except Exception:
        loader.rollback()
        raise
    finally:
        loader.close()
    summary["ok"] = True
    summary["mode"] = "dry-run" if dry_run else "committed"
    return summary


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="엑셀 데이터를 PostgreSQL에 적재합니다.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument("--file", "-f", help="엑셀 파일 경로 (.xlsx)")
    parser.add_argument("--sheet", "-s", help="적재할 시트 이름")
    parser.add_argument("--table", "-t", help="대상 테이블 (생략 시 시트 이름으로 추론)")
    parser.add_argument("--dry-run", action="store_true", help="트랜잭션을 롤백하여 미리보기만 수행")
    parser.add_argument("--create-template", metavar="PATH", help="한글 헤더 템플릿 엑셀을 생성")
    parser.add_argument(
        "--no-create-missing-municipalities",
        action="store_true",
        help="없는 지자체를 자동 생성하지 않음",
    )
    parser.add_argument("--json", action="store_true", help="결과를 JSON으로 출력")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if args.create_template:
        path = create_template(Path(args.create_template).expanduser().resolve())
        print(f"템플릿을 생성했습니다: {path}")
        return 0
    if not args.file:
        print("--file 또는 --create-template 이 필요합니다.", file=sys.stderr)
        return 2
    load_env_files()
    try:
        summary = load_file(
            Path(args.file).expanduser().resolve(),
            sheet=args.sheet,
            table=args.table,
            dry_run=args.dry_run,
            create_missing=not args.no_create_missing_municipalities,
        )
    except LoadError as exc:
        print(f"오류: {exc}", file=sys.stderr)
        return 1
    if args.json:
        print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
    else:
        print(f"파일: {summary['file']}")
        print(f"모드: {summary['mode']}")
        for sheet in summary["sheets"]:
            print(f"- {sheet['sheet']} → {sheet['table']}: {sheet['inserted']}건")
    return 0


if __name__ == "__main__":
    sys.exit(main())
